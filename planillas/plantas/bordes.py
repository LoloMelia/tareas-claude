# -*- coding: utf-8 -*-
"""Convencion de bordes del autor para todas las planillas.

Perimetro de cada tabla en linea llena; las lineas interiores, en puntos,
salvo las del encabezado, que van llenas.
"""
from openpyxl.styles import Border, Side

LLENA = Side(style="thin", color="000000")
PUNTOS = Side(style="dotted", color="000000")

def _tiene_borde(c):
    b = c.border
    return any(getattr(b, l).style for l in ("left", "right", "top", "bottom"))

def _es_encabezado(c):
    """Titulo = celda con relleno gris. La negrita sola no la hace titulo."""
    return bool(c.fill is not None and c.fill.patternType == "solid"
                and (c.fill.fgColor.rgb or "")[-6:].upper() not in ("FFFFFF", ""))

def aplicar(ws):
    """Reescribe los bordes de la hoja segun la convencion."""
    filas = ws.max_row; cols = ws.max_column
    marco = {}; cab = {}
    for f in ws.iter_rows():
        for c in f:
            if _tiene_borde(c):
                marco[(c.row, c.column)] = True
                cab[(c.row, c.column)] = _es_encabezado(c)
    for (r, co) in marco:
        c = ws.cell(row=r, column=co)
        lados = {}
        for lado, (dr, dc) in (("top", (-1, 0)), ("bottom", (1, 0)),
                               ("left", (0, -1)), ("right", (0, 1))):
            if not getattr(c.border, lado).style:
                lados[lado] = Side()
                continue
            vecino = (r + dr, co + dc)
            if vecino not in marco:                 # borde exterior de la tabla
                lados[lado] = LLENA
            elif cab[(r, co)] or cab[vecino]:       # toca un encabezado
                lados[lado] = LLENA
            else:
                lados[lado] = PUNTOS
        c.border = Border(**lados)
    return len(marco)
